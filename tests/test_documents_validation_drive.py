"""PDF/DOCX document ingest, SHACL validation endpoint and Drive backup-before-replace."""

from __future__ import annotations

import zipfile
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from sion_api.main import create_app
from sion_core import extras_report
from sion_core.optional import is_available
from sion_drive_store.store import DriveLayout
from sion_drive_store.upload import publish_to_mounted_drive
from sion_ingestion import document_ingest
from sion_ingestion.document_ingest import DocumentUnreadable, extract_text

needs_pypdf = pytest.mark.skipif(not is_available("pypdf"), reason="needs the 'documents' extra (pypdf)")
needs_docx = pytest.mark.skipif(not is_available("docx"), reason="needs the 'documents' extra (python-docx)")
needs_validation = pytest.mark.skipif(not extras_report()["validation"]["installed"], reason="needs the 'validation' extra")


def _client(tmp_path: Path) -> TestClient:
    return TestClient(create_app(database_url="sqlite://", auto_create_schema=True, ingest_roots=[tmp_path]))


# --------------------------------------------------------------------------- fixtures built in code


def minimal_pdf(lines: list[str]) -> bytes:
    """A one-page PDF with Helvetica text lines (ASCII), with a correct xref table."""
    ops = ["BT", "/F1 12 Tf", "72 720 Td", "14 TL"]
    for line in lines:
        escaped = line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        ops.append(f"({escaped}) Tj T*")
    ops.append("ET")
    stream = "\n".join(ops).encode("latin-1")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    out += b"".join(f"{offset:010d} 00000 n \n".encode() for offset in offsets)
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return bytes(out)


W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def minimal_docx(path: Path, paragraphs: list[tuple[str | None, str]], styles: dict[str, str]) -> Path:
    """Hand-written OOXML package; ``styles`` maps styleId -> style name (e.g. Korean Word uses ids like '1')."""
    def para(style, text):
        ppr = f'<w:pPr><w:pStyle w:val="{style}"/></w:pPr>' if style else ""
        return f"<w:p>{ppr}<w:r><w:t xml:space=\"preserve\">{text}</w:t></w:r></w:p>"

    body = "".join(para(s, t) for s, t in paragraphs)
    table = f"<w:tbl><w:tr><w:tc>{para(None, 'table cell (not a body paragraph)')}</w:tc></w:tr></w:tbl>"
    document = f'<?xml version="1.0" encoding="UTF-8"?><w:document xmlns:w="{W}"><w:body>{body}{table}</w:body></w:document>'
    style_xml = "".join(f'<w:style w:type="paragraph" w:styleId="{i}"><w:name w:val="{n}"/></w:style>' for i, n in styles.items())
    styles_doc = f'<?xml version="1.0" encoding="UTF-8"?><w:styles xmlns:w="{W}">{style_xml}</w:styles>'
    content_types = (
        '<?xml version="1.0" encoding="UTF-8"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
        '<Default Extension="xml" ContentType="application/xml"/>'
        '<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
        '<Override PartName="/word/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.styles+xml"/>'
        "</Types>"
    )
    rels = (
        '<?xml version="1.0" encoding="UTF-8"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>'
        "</Relationships>"
    )
    doc_rels = (
        '<?xml version="1.0" encoding="UTF-8"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" Target="styles.xml"/>'
        "</Relationships>"
    )
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", content_types)
        archive.writestr("_rels/.rels", rels)
        archive.writestr("word/document.xml", document)
        archive.writestr("word/_rels/document.xml.rels", doc_rels)
        archive.writestr("word/styles.xml", styles_doc)
    return path


PARAS = [("1", "펌프실 개요"), (None, "급수 펌프 2대"), ("Heading2", "Valve schedule"), (None, ""), (None, "V-1 gate valve")]
STYLES = {"1": "heading 1", "Heading2": "Heading 2"}
EXPECTED = "# 펌프실 개요\n급수 펌프 2대\n## Valve schedule\nV-1 gate valve"


# --------------------------------------------------------------------------- DOCX


def test_docx_builtin_reader_maps_heading_styles_and_skips_tables(tmp_path, monkeypatch):
    path = minimal_docx(tmp_path / "pump.docx", PARAS, STYLES)
    monkeypatch.setattr(document_ingest, "is_available", lambda module: False if module == "docx" else is_available(module))
    content = extract_text(path)
    assert content.extractor == "docx-builtin"
    assert content.text == EXPECTED


@needs_docx
def test_python_docx_and_builtin_reader_agree(tmp_path, monkeypatch):
    import docx

    path = tmp_path / "made.docx"
    document = docx.Document()
    document.add_heading("Pump room", level=1)
    document.add_paragraph("two supply pumps")
    document.add_heading("Valves", level=2)
    document.add_paragraph("V-1\tgate valve")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "in a table"
    document.save(str(path))
    with_library = extract_text(path)
    assert with_library.extractor == "python-docx"
    monkeypatch.setattr(document_ingest, "is_available", lambda module: False if module == "docx" else is_available(module))
    builtin = extract_text(path)
    assert builtin.text == with_library.text == "# Pump room\ntwo supply pumps\n## Valves\nV-1\tgate valve"


def test_broken_docx_is_unreadable(tmp_path):
    path = tmp_path / "broken.docx"
    path.write_bytes(b"not a zip")
    with pytest.raises(DocumentUnreadable):
        extract_text(path)


# --------------------------------------------------------------------------- PDF


@needs_pypdf
def test_pdf_text_extraction(tmp_path):
    path = tmp_path / "spec.pdf"
    path.write_bytes(minimal_pdf(["# Fire door schedule", "FD-1 60 min", "FD-2 30 min"]))
    content = extract_text(path)
    assert content.extractor == "pypdf" and content.page_count == 1
    assert [line.strip() for line in content.text.splitlines() if line.strip()] == ["# Fire door schedule", "FD-1 60 min", "FD-2 30 min"]


def test_pdf_without_pypdf_names_the_extra(tmp_path, monkeypatch):
    import sion_core.optional as optional

    path = tmp_path / "spec.pdf"
    path.write_bytes(minimal_pdf(["x"]))
    real_import = optional.importlib.import_module

    def fake_import(name, *args, **kwargs):
        if name == "pypdf":
            raise ImportError("blocked for test")
        return real_import(name, *args, **kwargs)

    monkeypatch.setattr(optional.importlib, "import_module", fake_import)
    with pytest.raises(ImportError, match=r"documents"):
        extract_text(path)


# --------------------------------------------------------------------------- API


def test_ingest_api_reports_skipped_files_instead_of_failing(tmp_path):
    note = tmp_path / "note.md"
    note.write_text("# Pump room\n", encoding="utf-8")
    docx_path = minimal_docx(tmp_path / "pump.docx", PARAS, STYLES)
    bad = tmp_path / "bad.docx"
    bad.write_bytes(b"garbage")
    empty_pdf = tmp_path / "scan.pdf"
    empty_pdf.write_bytes(minimal_pdf([]))
    other = tmp_path / "image.png"
    other.write_bytes(b"\x89PNG")
    with _client(tmp_path) as c:
        r = c.post("/api/v1/ingest/documents", json={"paths": [str(p) for p in (note, docx_path, bad, empty_pdf, other)]})
        assert r.status_code == 200, r.text
        body = r.json()
        assert body["files"] == [str(note), str(docx_path)]
        assert body["extractors"][str(note)] == "text"
        assert body["extractors"][str(docx_path)] in {"python-docx", "docx-builtin"}
        reasons = {Path(s["path"]).name: s["reason"] for s in body["skipped"]}
        assert set(reasons) == {"bad.docx", "scan.pdf", "image.png"}
        assert "DOCX could not be read" in reasons["bad.docx"]
        assert reasons["image.png"] == "unsupported type .png"
        if is_available("pypdf"):
            assert "no extractable text" in reasons["scan.pdf"]
        else:
            assert "documents" in reasons["scan.pdf"]
        # claims: 1 heading from the note + 2 headings from the DOCX
        assert body["evidence_created"] == 1 + 2
        assert c.post("/api/v1/ingest/documents", json={"paths": [str(note), str(docx_path)]}).json()["evidence_created"] == 0


@needs_pypdf
def test_ingest_api_pdf_round_trip(tmp_path):
    pdf = tmp_path / "spec.pdf"
    pdf.write_bytes(minimal_pdf(["# Fire door schedule", "# Hardware"]))
    with _client(tmp_path) as c:
        body = c.post("/api/v1/ingest/documents", json={"paths": [str(pdf)]}).json()
        assert body["files"] == [str(pdf)] and body["extractors"][str(pdf)] == "pypdf"
        assert body["evidence_created"] == 2
        doc = [e for e in c.get("/api/v1/entities", params={"limit": 50}).json() if e["entity_type_id"] == "Document"]
        assert doc and doc[0]["properties"]["page_count"] == 1


# --------------------------------------------------------------------------- SHACL


@needs_validation
def test_shacl_endpoint_conforms_for_api_written_graph(tmp_path):
    with _client(tmp_path) as c:
        assert c.get("/api/v1/validation/shacl").json()["conforms"] is True  # empty graph
        ids = [
            c.post("/api/v1/entities", json={"stable_key": k, "entity_type_id": "Concept", "name": k}).json()["id"]
            for k in ("a", "b")
        ]
        rel = c.post("/api/v1/relations", json={"stable_key": "a-b", "source_entity_id": ids[0], "target_entity_id": ids[1],
                                                 "relation_type_id": "RELATED_TO", "confidence": 0.7}).json()
        assert c.post("/api/v1/evidence", json={"relation_id": rel["id"], "source_uri": "urn:test"}).status_code == 201
        body = c.get("/api/v1/validation/shacl").json()
        assert body["conforms"] is True and body["violation_count"] == 0
        assert body["validated"] == {"entities": 2, "relations": 1, "evidence": 1} and body["truncated"] is False
        limited = c.get("/api/v1/validation/shacl", params={"limit": 1}).json()
        assert limited["truncated"] is True and limited["validated"]["entities"] == 1


@needs_validation
def test_shacl_reports_violations():
    import rdflib
    from rdflib import RDF, Literal, Namespace
    from sion_api import shacl
    from sion_api.config import load_settings

    sion = Namespace(shacl.SION_NS)
    graph = rdflib.Graph()
    graph.add((sion["relation/r1"], RDF.type, sion.Relation))
    graph.add((sion["relation/r1"], sion.source_id, Literal("a")))
    graph.add((sion["relation/r1"], sion.target_id, Literal("b")))
    graph.add((sion["relation/r1"], sion.relation_type, Literal("USES")))
    graph.add((sion["relation/r1"], sion.confidence, Literal(1.5)))
    graph.add((sion["evidence/e1"], RDF.type, sion.Evidence))
    graph.add((sion["evidence/e1"], sion.verification_state, Literal("bogus")))
    graph.add((sion["entity/x"], RDF.type, sion.Entity))  # no id, no name
    conforms, violations = shacl.validate_graph(graph, shacl.shapes_path(load_settings().ontology_path))
    assert conforms is False
    found = {(v["focus_node"].rsplit("/", 1)[-1], v["constraint"]) for v in violations}
    assert ("r1", "MaxInclusiveConstraintComponent") in found
    assert ("e1", "InConstraintComponent") in found and ("e1", "OrConstraintComponent") in found
    assert ("x", "MinCountConstraintComponent") in found


def test_shacl_endpoint_503_without_shapes_or_extra(tmp_path, monkeypatch):
    from sion_api import shacl

    monkeypatch.setenv(shacl.SHAPES_ENV, str(tmp_path / "missing.ttl"))
    with _client(tmp_path) as c:
        r = c.get("/api/v1/validation/shacl")
        assert r.status_code == 503
        assert ("missing.ttl" in r.json()["detail"]) or ("validation" in r.json()["detail"])
    monkeypatch.delenv(shacl.SHAPES_ENV)
    monkeypatch.setattr(shacl, "_libs", lambda: (_ for _ in ()).throw(shacl.ShaclUnavailable("install [validation]")))
    with _client(tmp_path) as c:
        assert c.get("/api/v1/validation/shacl").status_code == 503


# --------------------------------------------------------------------------- Drive publish


def test_publish_backs_up_changed_files_and_never_deletes(tmp_path):
    stage = tmp_path / "stage"
    drive = tmp_path / "drive"
    (stage / "exports").mkdir(parents=True)
    (stage / "exports" / "graph.json").write_text('{"v": 1}', encoding="utf-8")
    (stage / "notes.md").write_text("same", encoding="utf-8")

    first = publish_to_mounted_drive(stage, drive)
    destination = drive / DriveLayout().project_root
    assert first["copied"] == 2 and first["archived"] == 0 and first["history"] is None
    (destination / "drive-only.txt").write_text("keep me", encoding="utf-8")

    (stage / "exports" / "graph.json").write_text('{"v": 2}', encoding="utf-8")  # same size, new content
    second = publish_to_mounted_drive(stage, drive)
    assert second["copied"] == 1 and second["unchanged"] == 1 and second["archived_paths"] == ["exports/graph.json"]
    history = Path(second["history"])
    assert history.parent == destination / ".history"
    assert (history / "exports" / "graph.json").read_text(encoding="utf-8") == '{"v": 1}'
    assert (destination / "exports" / "graph.json").read_text(encoding="utf-8") == '{"v": 2}'
    assert (destination / "drive-only.txt").read_text(encoding="utf-8") == "keep me"

    third = publish_to_mounted_drive(stage, drive)
    assert third["copied"] == 0 and third["archived"] == 0 and third["unchanged"] == 2


def test_publish_without_drive_root(tmp_path, monkeypatch):
    import sion_drive_store.upload as upload

    monkeypatch.setattr(upload, "detect_google_drive_root", lambda: None)
    assert publish_to_mounted_drive(tmp_path) == {"published": False, "reason": "google drive root not mounted"}
